#!/usr/bin/env python3
"""Convert the ZeonAI Agent Inventory workbook into normalized data.json.

Usage:
    python build/convert.py [path/to/workbook.xlsx] [-o data/data.json]

Reads only the canonical sheets (Master Data, Agent Inventory Live,
Agents Backlog, Summary Dashboard). The per-category narrative sheets and the
"Agent Inv (Dharm)" working copy are intentionally ignored -- their content is
already represented in the canonical tables above.
"""
import argparse
import datetime as dt
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

DEFAULT_XLSX = "inventory/Zeon_AI_Agents_Inventory.xlsx"
DEFAULT_OUT = "data/data.json"

KNOWN_PRODUCT_LINES = {"DeepDelve", "DeepProbe", "Cross-Platform"}

# Product line for categories that are not present in the Master Data table
# (the source sheet only maps CAT-01..CAT-09).
PRODUCT_LINE_FALLBACK = {
    "CAT-10": "Cross-Platform",  # Construction & Field Operations
    "CAT-11": "Cross-Platform",  # Procurement & Supply Chain
}

LIVE_HINTS = ("user", "client", "document", "progress", "live")
PLANNED_HINTS = ("to be deployed", "in development", "tbd", "planned")


def s(v):
    """Stringify a cell, trimming whitespace; None -> ''."""
    if v is None:
        return ""
    return str(v).strip()


def normalize_metric_value(raw):
    """Return (label, pct) for a Key Metric Value cell.

    Numeric fractions (0.85) -> ("85%", 85). Strings ("98%+", "100,000+")
    are kept verbatim with pct=None.
    """
    if raw is None:
        return "", None
    if isinstance(raw, (int, float)):
        if 0 < raw <= 1:
            pct = round(raw * 100)
            return f"{pct}%", pct
        return str(raw), None
    text = str(raw).strip()
    return text, None


def deployment_status(detail_blob):
    """Classify a category/agent as Live, In Progress, or Planned."""
    t = detail_blob.lower()
    if any(h in t for h in PLANNED_HINTS):
        return "Planned"
    if "progress" in t or "development" in t:
        return "In Progress"
    if any(h in t for h in LIVE_HINTS):
        return "Live"
    return "Planned"


def parse_master_data(ws):
    """Return (product_lines, industries, deployments_by_cat)."""
    product_lines = {}            # CategoryID -> product line
    industries = {}               # CategoryID -> [industry, ...]
    deployments = {}              # CategoryID -> [{client, details}, ...]
    section = None
    for row in ws.iter_rows(values_only=True):
        c0, c1, c2, c3 = (s(row[i]) if i < len(row) else "" for i in range(4))
        if c0.startswith("TABLE 1"):
            section = "cat"
            continue
        if c0.startswith("TABLE 4"):
            section = "dep"
            continue
        if c0 in ("CategoryID", "DeploymentID") or not c0:
            continue
        if section == "cat" and c0.startswith("CAT-"):
            if c2 in KNOWN_PRODUCT_LINES:
                product_lines.setdefault(c0, c2)
            elif c2:
                industries.setdefault(c0, [])
                if c2 not in industries[c0]:
                    industries[c0].append(c2)
        elif section == "dep" and c0.startswith("DEP-"):
            deployments.setdefault(c1, []).append(
                {"client": c2, "details": c3}
            )
    return product_lines, industries, deployments


def parse_agents(ws):
    agents = []
    cat_names = {}  # CategoryID -> Category Name (from inventory)
    rows = list(ws.iter_rows(values_only=True))
    # Header is on the 2nd row; data starts on the 3rd.
    for row in rows[2:]:
        aid = s(row[0])
        if not aid or not aid.upper().startswith("AGT-"):
            continue
        cat_id = s(row[2])
        cat_name = s(row[3])
        if cat_id:
            cat_names.setdefault(cat_id, cat_name)
        label, pct = normalize_metric_value(row[9] if len(row) > 9 else None)
        agents.append({
            "id": aid,
            "name": s(row[1]),
            "categoryId": cat_id,
            "categoryName": cat_name,
            "businessFunction": s(row[4]),
            "problem": s(row[5]),
            "before": s(row[6]),
            "after": s(row[7]),
            "metric": s(row[8]),
            "metricValue": label,
            "metricPct": pct,
            "description": s(row[10]) if len(row) > 10 else "",
            "internalComment": s(row[11]) if len(row) > 11 else "",
        })
    return agents, cat_names


def parse_backlog(ws):
    backlog = []
    rows = list(ws.iter_rows(values_only=True))
    for row in rows[1:]:  # header on row 1
        num = s(row[0])
        name = s(row[1])
        if not name:
            continue
        backlog.append({
            "n": num,
            "name": name,
            "description": s(row[2]),
            "category": s(row[3]),
            "erpFlag": s(row[4]).lower() in ("yes", "true", "y"),
            "internalComment": s(row[5]) if len(row) > 5 else "",
            "sprint": s(row[6]) if len(row) > 6 else "",
        })
    return backlog


def parse_summary(ws):
    """Extract KPIs, key sectors, ROI metrics and differentiators."""
    kpis, sectors, roi, diffs = [], [], [], []
    section = None
    for row in ws.iter_rows(values_only=True):
        cells = [s(c) for c in row]
        c0 = cells[0] if cells else ""
        c1 = cells[1] if len(cells) > 1 else ""
        c3 = cells[3] if len(cells) > 3 else ""
        head = c0.upper()
        if head.startswith("DEPLOYMENT SUMMARY"):
            section = "deploy"; continue
        if head.startswith("TYPICAL ROI"):
            section = "roi"; continue
        if head.startswith("KEY DIFFERENTIATORS"):
            section = "diff"; continue
        if c0.startswith("Contact"):
            section = None; continue
        if section == "deploy":
            if c0 and c0 != "Metric" and c1:
                kpis.append({"label": c0, "value": c1})
            if c3 and c3 != "Key Sectors":
                sectors.append(c3)
        elif section == "roi":
            if c0 and c0 != "Efficiency Metric" and c1:
                roi.append({"metric": c0, "improvement": c1})
        elif section == "diff":
            if c0 and c0 != "Differentiator" and c1:
                diffs.append({"title": c0, "details": c1})
    return kpis, sectors, roi, diffs


def build(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    product_lines, industries, deployments = parse_master_data(wb["Master Data"])
    agents, inv_cat_names = parse_agents(wb["Agent Inventory Live"])
    backlog = parse_backlog(wb["Agents Backlog"])
    kpis, sectors, roi, diffs = parse_summary(wb["Summary Dashboard"])

    # Category universe = everything referenced anywhere.
    cat_ids = set(inv_cat_names) | set(product_lines) | set(industries) | set(deployments)
    agents_by_cat = {}
    for a in agents:
        agents_by_cat.setdefault(a["categoryId"], 0)
        agents_by_cat[a["categoryId"]] += 1

    categories = []
    for cid in sorted(cat_ids):
        deps = deployments.get(cid, [])
        status = deployment_status(" ".join(d["client"] + " " + d["details"] for d in deps)) if deps else "Planned"
        pl = product_lines.get(cid) or PRODUCT_LINE_FALLBACK.get(cid, "Cross-Platform")
        categories.append({
            "id": cid,
            "name": inv_cat_names.get(cid, cid),
            "productLine": pl,
            "industries": industries.get(cid, []),
            "deployments": deps,
            "agentCount": agents_by_cat.get(cid, 0),
            "status": status,
        })

    # Tag each agent with its category's deployment status & product line.
    cat_lookup = {c["id"]: c for c in categories}
    for a in agents:
        c = cat_lookup.get(a["categoryId"])
        a["status"] = c["status"] if c else "Planned"
        a["productLine"] = c["productLine"] if c else "Cross-Platform"

    return {
        "meta": {
            "title": "ZeonAI Agent Portfolio",
            "subtitle": "DeepDelve · DeepProbe — Enterprise Agentic AI",
            "contact": "AlphaData Sales Team · www.zeonai.com",
            "generatedAt": dt.datetime.now().strftime("%d %b %Y"),
            "sourceFile": os.path.basename(xlsx_path),
            "kpis": kpis,
            "keySectors": sectors,
            "roi": roi,
            "differentiators": diffs,
        },
        "categories": categories,
        "agents": agents,
        "backlog": backlog,
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("xlsx", nargs="?", default=DEFAULT_XLSX)
    ap.add_argument("-o", "--out", default=DEFAULT_OUT)
    args = ap.parse_args()

    data = build(args.xlsx)
    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    m = data["meta"]
    print(f"Wrote {args.out}")
    print(f"  agents={len(data['agents'])} categories={len(data['categories'])} "
          f"backlog={len(data['backlog'])}")
    print(f"  kpis={len(m['kpis'])} roi={len(m['roi'])} differentiators={len(m['differentiators'])} "
          f"sectors={len(m['keySectors'])}")
    deps = sum(len(c['deployments']) for c in data['categories'])
    print(f"  deployments={deps}")


if __name__ == "__main__":
    main()
